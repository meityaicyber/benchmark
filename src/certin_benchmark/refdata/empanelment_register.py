"""Conversion of CERT-In's register of empanelled auditing organisations (a spreadsheet) into a dated register.

Besides the list it publishes (Empanel_org.pdf, converted by empanelment_pdf.py), CERT-In supplies a spreadsheet of
the empanelled organisations giving, for each:

* its CERT-In reference ("Code(Reference)"): the round of empanelment and a number, in forms such as D<n>-<year>,
  RE-<year>-<n>, E-<n>, F-<n>, F-<year>-<n> and G-<n>;
* its AMBAK code ("AMBAK Code (Primary)"): the organisation's identifier on AMBAK (Audit Monitoring, Benchmarking,
  Analysis and Kinetic Interventions), the platform CERT-In launched at CERT-In SAMVAAD 2026, ten digits read as a
  year, a month and a serial number;
* its name, with any former name ("Formerly ..."); and
* the date up to which its empanelment is valid ("Empanelment Validity").

Neither the codes nor the validity appear in the published list. The spreadsheet is the "validity of empanelment"
of the CERT-In empanelment records (Suite Section 11.1), obtained from CERT-In directly as Section 11.2 requires.

The converter writes staging/empanelment_register/register_<date>.json, with the spreadsheet beside it:

    {"register_date": "YYYY-MM-DD", "source": {...},
     "organisations": [{"row_no", "serial", "certin_reference", "ambak_code", "name", "former_names",
                        "other_names", "valid_to"}],
     "warnings": [...]}

row_no is the row of the sheet, serial the sheet's own S.No. build.py keeps every register received under its date,
as it keeps the empanelment sets, and links each row to the organisation of each set by the organisation's name, a
former or other name the register gives, or the AMBAK code of a row already linked. The register in force on a
report's release date is the latest dated on or before it, else the earliest (store.py).

Reading the spreadsheet needs openpyxl, on the preparation host only; the identifier and name helpers are used
inside the environment as well.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from .empanelment_pdf import norm_org, org_name_variants

# India Standard Time, in which a register without --register-date is dated by the time it was last saved
IST_OFFSET = _dt.timedelta(hours=5, minutes=30)

# Column headings, matched case-insensitively against each cell of the heading row (first field that matches)
HEADINGS = {
    "serial": r"^(s\.?\s*no|sr\.?\s*no|sl\.?\s*no|serial)\b",
    "certin_reference": r"code\s*\(?\s*reference|reference\s*(code|no|number)|empanelment\s*(no|number|code|id)\b",
    "ambak_code": r"\bambak\b",
    "name": r"organi[sz]ation('?s)?\s*name|name\s+of\s+(the\s+)?(empanelled\s+)?(auditing\s+)?organi[sz]ation|"
            r"^(empanelled\s+)?organi[sz]ation$",
    "valid_to": r"validity|valid\s*(up\s*to|upto|till|until)",
}

# Words a report may set before an identifier: 'AMBAK Code: ...', 'CERT-In Empanelment No. ...', 'CERT-In/...'
_ID_PREFIX = re.compile(r"^\s*(?:ambak(?:\s*(?:code|id|no\.?|number))?|cert\s*-?\s*in(?:\s*(?:empanelment|reference|ref\.?))?"
                        r"(?:\s*(?:code|id|no\.?|number))?)\s*[:/#.\-]*\s*", re.I)

# Former and other names. In brackets anywhere after the name ('(Formerly X)', '(Formerly known as X (India) Ltd)',
# '(Now X)'); without brackets only after the name and only as 'formerly X' or '<previously|earlier> known as X', so
# that a name beginning 'Earlier' or 'Previously' is not split.
_INNER = r"((?:[^()]|\([^()]*\))+?)"
_FORMER_WORDS = r"(?:formerly|previously|earlier)\b(?:\s+known\s+as\b|\s+called\b)?"
_OTHER_WORDS = r"(?:now\b(?:\s+known\s+as\b)?|also\s+known\s+as\b|a\.?\s?k\.?\s?a\b\.?)"
PAREN_FORMER = re.compile(r"\(\s*" + _FORMER_WORDS + r"\s*[:\-\u2013]?\s*" + _INNER + r"\s*(?:\)|$)", re.I)
PAREN_OTHER = re.compile(r"\(\s*" + _OTHER_WORDS + r"\s*[:\-\u2013]?\s*" + _INNER + r"\s*(?:\)|$)", re.I)
BARE_FORMER = re.compile(r"(?<=[\s,;\-\u2013])(?:formerly\b(?:\s+known\s+as\b)?|(?:previously|earlier)\s+known\s+as\b)"
                         r"\s*[:\-\u2013]?\s*" + _INNER + r"\s*\)?\s*$", re.I)
BARE_OTHER = re.compile(r"(?<=[\s,;\-\u2013])(?:now|also)\s+known\s+as\b\s*[:\-\u2013]?\s*" + _INNER + r"\s*\)?\s*$",
                        re.I)


def strip_id_prefix(value) -> str:
    return _ID_PREFIX.sub("", str(value or "")).strip()


def norm_reference(value) -> str:
    """A CERT-In reference for comparison: its letters and numbers in order, ignoring case, the spaces and
    punctuation between them and leading zeros ('RE-2099-07', 're 2099 7' and 'CERT-In/RE-2099-07' are the same
    reference; 'F-2099-002' and 'F2099002' are not)."""
    groups = re.findall(r"[A-Za-z]+|\d+", strip_id_prefix(value))
    return "-".join(g.upper() if g.isalpha() else str(int(g)) for g in groups)


def norm_ambak(value) -> str | None:
    """An AMBAK code as a string of digits, from a number or a text cell ('AMBAK Code: ...' accepted); None where the
    value is not one number of at least six digits."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    if isinstance(value, int):
        s = str(value)
    else:
        s = re.sub(r"\.0+$", "", strip_id_prefix(value))           # a number written out from a float
        if re.fullmatch(r"\d{1,3}(?:,\d{3})+", s):                   # thousands separators
            s = s.replace(",", "")
    return s if re.fullmatch(r"\d{6,}", s) else None


def _names_in(name: str, paren: re.Pattern, bare: re.Pattern) -> tuple[list[str], str]:
    out = []
    rest = str(name or "")
    for m in paren.finditer(rest):
        out.append(m.group(1))
    rest = paren.sub(" ", rest)
    m = bare.search(rest)
    if m and m.start() > 0:
        out.append(m.group(1))
        rest = rest[:m.start()]
    cleaned = [re.sub(r"\s+", " ", n).strip(" .,;:-–") for n in out]
    return [n for n in cleaned if n and norm_org(n)], rest


def former_names(name: str) -> list[str]:
    """Former names a register gives with the current one: 'Astra Example Pvt. Ltd. (Formerly Old Example Advisors)'
    gives 'Old Example Advisors'."""
    return _names_in(name, PAREN_FORMER, BARE_FORMER)[0]


def other_names(name: str) -> list[str]:
    """Other current names a register gives: '(Now X)', 'now known as X', 'also known as X'."""
    return _names_in(name, PAREN_OTHER, BARE_OTHER)[0]


def name_variants(name: str, former: list[str] | None = None, other: list[str] | None = None) -> list[str]:
    """Every normalised form under which a register row names its organisation: the name as written, its current
    part, and each former or other name (taken from the name where not given)."""
    out = list(org_name_variants(name))
    for f in (former if former is not None else former_names(name)):
        out.extend(org_name_variants(f))
    for f in (other if other is not None else other_names(name)):
        out.extend(org_name_variants(f))
    return [v for v in dict.fromkeys(out) if v]


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y", "%d.%m.%y", "%d-%b-%y", "%d-%b-%Y",
                "%d %b %Y", "%d %B %Y", "%d-%B-%Y", "%d %b %y", "%b %d %Y", "%B %d %Y")


def _date(value) -> str | None:
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # an Excel serial number (days since 30 December 1899), in the range of dates from 1954 to 2119
        if 20000 <= value <= 80000:
            return (_dt.date(1899, 12, 30) + _dt.timedelta(days=int(value))).isoformat()
        return None
    s = _clean(value)
    if not s:
        return None
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", s)
    if m:
        try:
            return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return None
    s = re.sub(r"^(?:valid\s*)?(?:till|upto|up\s*to|until|w\.?e\.?f\.?)\s*[:\-]?\s*", "", s, flags=re.I)
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s, flags=re.I).replace(",", " ")
    s = re.sub(r"\s+", " ", s).strip()
    for fmt in DATE_FORMATS:
        try:
            return _dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _heading_map(row) -> dict[str, int]:
    """Columns of each field. Where several columns name the AMBAK code, the one marked 'Primary' is read."""
    found: dict[str, list[int]] = {}
    for i, cell in enumerate(row):
        text = _clean(cell).lower()
        if not text:
            continue
        for field, rx in HEADINGS.items():
            if re.search(rx, text):
                found.setdefault(field, []).append(i)
                break
    cols = {f: v[0] for f, v in found.items()}
    if len(found.get("ambak_code", [])) > 1:
        primary = [i for i in found["ambak_code"] if "primary" in _clean(row[i]).lower()]
        cols["ambak_code"] = primary[0] if primary else found["ambak_code"][0]
    return cols


def _grid(ws) -> list[list]:
    """The sheet's values, each merged range filled with the value of its first cell."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in getattr(ws, "merged_cells", None).ranges if getattr(ws, "merged_cells", None) else []:
        top = rows[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(rows) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
            for c in range(rng.min_col - 1, rng.max_col):
                if c < len(rows[r]) and rows[r][c] is None:
                    rows[r][c] = top
    return rows


def check_date(value: str | None) -> str | None:
    if value is None:
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value)):
        raise ValueError(f"date {value!r} is not of the form YYYY-MM-DD")
    _dt.date.fromisoformat(str(value))
    return str(value)


def read_register(path: str | Path, register_date: str | None = None) -> dict:
    """Read the register from the first sheet whose heading row names the organisation and an identifier."""
    import openpyxl  # preparation host only

    register_date = check_date(register_date)
    path = Path(path)
    data = path.read_bytes()
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        props = wb.properties
        modified = props.modified or props.created
        # The workbook records when it was last saved in UTC. The register is dated in India Standard Time, as the
        # list's PDF is dated by its own local date, so that a sheet saved between midnight and 05:30 IST keeps that
        # day's date.
        document_date = (modified + IST_OFFSET).date().isoformat() if modified else None
        found = None
        others = []
        for ws in wb.worksheets:
            rows = _grid(ws)
            for h, row in enumerate(rows[:15]):
                cols = _heading_map(row)
                if "name" in cols and ("ambak_code" in cols or "certin_reference" in cols):
                    if found is None:
                        found = (ws.title, h, rows, cols)
                    else:
                        others.append(ws.title)
                    break
        if found is not None:
            out = _convert(path, data, *found, document_date, register_date)
            out["warnings"].extend(f"sheet {t!r} also has a heading row of a register and was not read; only sheet "
                                   f"{found[0]!r} is" for t in others)
            return out
    finally:
        wb.close()
    raise ValueError(f"{path.name}: no sheet with a heading row naming the organisation and an AMBAK code or "
                     "CERT-In reference")


def _convert(path: Path, data: bytes, sheet: str, h: int, rows: list, cols: dict, document_date: str | None,
             register_date: str | None) -> dict:
    def cell(row, field):
        i = cols.get(field)
        return row[i] if i is not None and i < len(row) else None

    organisations, warnings = [], []
    for k in range(h + 1, len(rows)):
        row = rows[k]
        if not any(_clean(v) for v in row):
            continue
        sheet_row = k + 1
        name = _clean(cell(row, "name"))
        serial = cell(row, "serial")
        serial = int(serial) if isinstance(serial, (int, float)) and not isinstance(serial, bool) else (_clean(serial) or None)
        ambak_raw, ref_raw, valid_raw = cell(row, "ambak_code"), cell(row, "certin_reference"), cell(row, "valid_to")
        ambak = norm_ambak(ambak_raw)
        ref = _clean(ref_raw) or None
        valid_to = _date(valid_raw)
        label = f"row {sheet_row} ({name or 'no name'})"
        if not name:
            warnings.append(f"{label}: no organisation name; row not read")
            continue
        if "ambak_code" in cols and ambak is None:
            warnings.append(f"{label}: AMBAK code {_clean(ambak_raw)!r} is not one number" if _clean(ambak_raw)
                            else f"{label}: no AMBAK code")
        if "certin_reference" in cols and not ref:
            warnings.append(f"{label}: no CERT-In reference")
        if "valid_to" in cols and valid_to is None:
            warnings.append(f"{label}: validity {_clean(valid_raw)!r} is not a date" if _clean(valid_raw)
                            else f"{label}: no validity date")
        organisations.append({"row_no": sheet_row, "serial": serial, "certin_reference": ref, "ambak_code": ambak,
                              "name": name, "former_names": former_names(name), "other_names": other_names(name),
                              "valid_to": valid_to})
    if not organisations:
        raise ValueError(f"{path.name}: no organisations read from sheet {sheet!r}")
    warnings.extend(_consistency_warnings(organisations))
    created = register_date or document_date or _dt.date.fromtimestamp(path.stat().st_mtime).isoformat()
    return {
        "register_date": created,
        "source": {"publisher": "CERT-In", "title": "Empanelled organisations with identifiers and validity",
                   "file": path.name, "sha256": hashlib.sha256(data).hexdigest(), "sheet": sheet,
                   "document_date": document_date, "rows": len(organisations),
                   "headings": {f: _clean(rows[h][i]) for f, i in cols.items()},
                   "converter": "certin_benchmark.refdata.empanelment_register"},
        "organisations": organisations,
        "warnings": warnings,
    }


def _consistency_warnings(organisations: list[dict]) -> list[str]:
    out = []
    by_ambak: dict[str, list[str]] = {}
    by_ref: dict[str, list[str]] = {}
    by_name: dict[str, list[str]] = {}
    for o in organisations:
        if o["ambak_code"]:
            by_ambak.setdefault(o["ambak_code"], []).append(o["name"])
        if o["certin_reference"]:
            by_ref.setdefault(norm_reference(o["certin_reference"]), []).append(o["name"])
        by_name.setdefault(norm_org(o["name"]), []).append(o["name"])
    for code, names in by_ambak.items():
        if len(names) > 1:
            out.append(f"AMBAK code {code} is given to {len(names)} rows: {'; '.join(names)}")
    for ref, names in by_ref.items():
        if len(names) > 1:
            out.append(f"CERT-In reference {ref} is given to {len(names)} rows: {'; '.join(names)}; it identifies "
                       "none of them")
    for names in by_name.values():
        if len(names) > 1:
            out.append(f"organisation listed more than once: {'; '.join(names)}")
    # The codes read as year, month and serial number: a code of another length than most, or whose month is not 01
    # to 12, is noted
    codes = [o["ambak_code"] for o in organisations if o["ambak_code"]]
    usual = Counter(len(c) for c in codes).most_common(1)[0][0] if codes else None
    for o in organisations:
        code = o["ambak_code"]
        if not code:
            continue
        if len(code) != usual or (len(code) == 10 and not 1 <= int(code[4:6]) <= 12):
            out.append(f"AMBAK code {code} of {o['name']} does not have the form of the other codes (year, month and "
                       "serial number)")
    return out


def write_register(xlsx_path: str | Path, out_dir: str | Path, register_date: str | None = None) -> Path:
    data = read_register(xlsx_path, register_date)
    out = Path(out_dir) / f"register_{data['register_date']}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        raise FileExistsError(f"{out} exists; a register is never overwritten")
    out.write_text(json.dumps(data, indent=1, ensure_ascii=False))
    # CERT-In's original is kept with the register, so that the package retains it (Section 11.2, step 3)
    original = out.with_suffix(Path(xlsx_path).suffix.lower() or ".xlsx")
    if not original.exists():
        original.write_bytes(Path(xlsx_path).read_bytes())
    return out
